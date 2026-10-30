#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
採購案件流程狀態機 procure.py（臺北市政府公務人員訓練處 總務組 採購承辦用）

由 Claude Code 的 /bid-run 技能呼叫，負責「確定性」的工作：
  - 建立案件（以案號為鍵）、依招標方式載入階段序列
  - 追蹤每個階段狀態：pending / in_progress / awaiting_confirm / confirmed / skipped
  - 記錄承辦人的確認、修改意見、跳過理由；同步產生人類可讀的 STATE.md
  - 履約期間插入「契約變更」「後續擴充」「分期驗收」臨時階段
  - 判斷進行到一半才確定的招標方式時，重排尚未開始的階段（replan）
  - 把各階段 Markdown 轉成 Word（export）

Claude 負責「非確定性」的工作：讀業務單位檔案、讀舊檔、擬稿、核對。
分工與鐵則見 CLAUDE.md 與 .claude/skills/bid-run/SKILL.md。

只用 Python 3 標準函式庫；export 時若有 python-docx 或 pandoc 會自動使用。
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
CASES_DIR = ROOT / "cases"
TEMPLATES_DIR = ROOT / "templates"
STEPS_DIR = ROOT / "steps"
REFERENCE_DIR = ROOT / "references"
CONFIG_FILE = ROOT / "config.json"
ACTIVE_FILE = CASES_DIR / ".active"

STATUS_LABEL = {
    "pending": "未開始",
    "in_progress": "作業中",
    "awaiting_confirm": "待承辦人確認",
    "confirmed": "已確認",
    "skipped": "已跳過",
}

# ---------------------------------------------------------------------------
# 階段目錄
#   folder    ：cases/<案號>/ 下的階段資料夾
#   step      ：steps/ 下給 Claude 的作業指示檔
#   templates ：templates/ 內的範本（開始階段時複製進資料夾，已存在不覆蓋）
#   confirm   ：這一步結束時要請承辦人確認什麼
# ---------------------------------------------------------------------------
STAGE_CATALOG: dict[str, dict] = {
    "intake": {
        "name": "收件", "folder": "01-intake", "step": "01-intake.md",
        "templates": ["01-summary"],
        "desc": "讀業務單位簽與附件，轉出文字，列摘要、缺漏與疑點。",
        "confirm": "摘要是否正確、缺漏由誰補",
    },
    "judge": {
        "name": "判斷", "folder": "02-judgement", "step": "02-judge.md",
        "templates": ["02-judgement"],
        "desc": "判斷金額級距、招標與決標方式、代辦、等標期、保證金、評選委員會。",
        "confirm": "各項判斷與法條；確認後據以決定後續階段（replan）",
    },
    "memo": {
        "name": "招標文件核定簽呈", "folder": "03-approval-memo", "step": "03-memo.md",
        "templates": ["03-approval-memo", "03-diff"],
        "desc": "依範本與舊檔撰寫招標文件核定簽呈（txt＋docx）及與舊檔差異對照表。",
        "confirm": "簽呈全文逐段確認",
    },
    "crosscheck": {
        "name": "交叉核對", "folder": "04-crosscheck", "step": "04-crosscheck.md",
        "templates": ["04-crosscheck"],
        "desc": "簽呈、招標資訊一覽表、公告稿、投標須知、契約、評選須知、標單互相核對。",
        "confirm": "每一項不一致的處理方式",
    },
    "announce_fields": {
        "name": "公告欄位", "folder": "05-announce-fields", "step": "05-announce.md",
        "templates": ["05-announce-fields", "05-schedule"],
        "desc": "整理政府電子採購網公告欄位與日程試算（由承辦人自行上網登打）。",
        "confirm": "欄位內容與日期",
    },
    "archive_tender": {
        "name": "招標文件歸檔", "folder": "06-archive", "step": "06-archive.md",
        "templates": ["06-INDEX"],
        "desc": "統一檔名、建立索引、列出承辦人自行處理的待辦（用印、上簽、公告、委員登載）。",
        "confirm": "歸檔清單與待辦",
    },
    "announce": {
        "name": "公告與等標期", "folder": "07-announce", "step": "07-announce.md",
        "templates": ["07-announce-log"],
        "desc": "記錄公告日、等標期、領標、釋疑與更正公告；核對等標期是否需延長。",
        "confirm": "公告紀錄；是否有釋疑或更正需另簽",
    },
    "bid_opening": {
        "name": "開標", "folder": "08-opening", "step": "08-opening.md",
        "templates": ["08-opening-record", "08-qualification", "08-opening-memo"],
        "desc": "擬開標紀錄、資格審查表與開標結果簽呈（流標／進入評選／當場決標）。",
        "confirm": "開標紀錄與開標結果簽呈",
    },
    "evaluation": {
        "name": "評選／評審", "folder": "09-evaluation", "step": "09-evaluation.md",
        "templates": ["09-committee-memo", "09-evaluation-record", "09-evaluation-result-memo"],
        "desc": "評選委員會成立簽、評選會議紀錄、評選結果（優勝序位）簽呈。",
        "confirm": "評選會議紀錄與評選結果簽呈",
    },
    "negotiation": {
        "name": "議價", "folder": "10-negotiation", "step": "10-negotiation.md",
        "templates": ["10-reserve-price-memo", "10-negotiation-record"],
        "desc": "底價核定簽（訂底價者）與議價／比減價格紀錄。",
        "confirm": "底價簽與議價紀錄",
    },
    "award": {
        "name": "決標", "folder": "11-award", "step": "11-award.md",
        "templates": ["11-award-memo", "11-award-record", "11-award-announce-fields"],
        "desc": "決標簽呈、決標紀錄與決標公告欄位。",
        "confirm": "決標簽呈與決標公告欄位",
    },
    "contract": {
        "name": "簽約", "folder": "12-contract", "step": "12-contract.md",
        "templates": ["12-contract-memo", "12-contract-fill", "12-bond-check"],
        "desc": "簽約簽呈、契約書填寫內容對照、履約保證金核對。",
        "confirm": "簽約簽呈與契約填寫內容",
    },
    "performance": {
        "name": "履約管理", "folder": "13-performance", "step": "13-performance.md",
        "templates": ["13-performance-log"],
        "desc": "開工、進度、監造、物調等履約追蹤；契約變更與後續擴充以臨時階段插入。",
        "confirm": "履約紀錄；是否需插入契約變更或後續擴充",
    },
    "acceptance": {
        "name": "驗收", "folder": "14-acceptance", "step": "14-acceptance.md",
        "templates": ["14-acceptance-record", "14-acceptance-memo"],
        "desc": "驗收（初驗、正式驗收）紀錄與驗收簽呈，含減價收受或限期改善。",
        "confirm": "驗收紀錄與驗收簽呈",
    },
    "closeout": {
        "name": "結案歸檔", "folder": "15-closeout", "step": "15-closeout.md",
        "templates": ["15-closeout-memo", "15-INDEX"],
        "desc": "結案付款簽、保證金退還、保固起算、全案檔案索引。",
        "confirm": "結案簽呈與全案索引",
    },
    # ---- 履約期間可插入的臨時階段 ----
    "change": {
        "name": "契約變更", "folder": "A1-change", "step": "A1-change.md",
        "templates": ["A1-change-memo", "A1-change-agreement"],
        "desc": "契約變更簽呈與變更協議書（數量、工期、金額、規格）。",
        "confirm": "契約變更簽呈與協議書", "adhoc": True,
    },
    "expansion": {
        "name": "後續擴充", "folder": "A2-expansion", "step": "A2-expansion.md",
        "templates": ["A2-expansion-memo", "10-negotiation-record"],
        "desc": "依採購法第22條第1項第7款辦理後續擴充之簽呈與議價紀錄。",
        "confirm": "後續擴充簽呈與議價紀錄", "adhoc": True,
    },
    "partial_acceptance": {
        "name": "分期驗收", "folder": "A3-partial-acceptance", "step": "14-acceptance.md",
        "templates": ["14-acceptance-record", "14-acceptance-memo"],
        "desc": "分期（部分）驗收紀錄與簽呈。",
        "confirm": "分期驗收紀錄與簽呈", "adhoc": True,
    },
}

ADHOC_ALIASES = {
    "契變": "change", "契約變更": "change", "change": "change", "變更": "change",
    "後續擴充": "expansion", "擴充": "expansion", "expansion": "expansion",
    "分期驗收": "partial_acceptance", "部分驗收": "partial_acceptance", "partial": "partial_acceptance",
}

PRE_TENDER = ["intake", "judge", "memo", "crosscheck", "announce_fields", "archive_tender"]
POST_AWARD = ["contract", "performance", "acceptance", "closeout"]

PROFILES: dict[str, dict] = {
    "公開招標_最有利標": {
        "desc": "公開招標、最有利標（本處工程案常用；需評選委員會，不訂底價）",
        "stages": PRE_TENDER + ["announce", "bid_opening", "evaluation", "award"] + POST_AWARD,
    },
    "公開招標_最低標": {
        "desc": "公開招標、最低標（訂底價，開標後審標、比減價格、當場決標）",
        "stages": PRE_TENDER + ["announce", "bid_opening", "negotiation", "award"] + POST_AWARD,
    },
    "限制性招標_公開評選": {
        "desc": "限制性招標、公開評選（採購法第22條第1項第9、10款等；評選後議價）",
        "stages": PRE_TENDER + ["announce", "bid_opening", "evaluation", "negotiation", "award"] + POST_AWARD,
    },
    "限制性招標_議價": {
        "desc": "限制性招標、不經公告直接議價（採購法第22條第1項各款，簽呈須敘明理由）",
        "stages": ["intake", "judge", "memo", "crosscheck", "archive_tender", "negotiation", "award"] + POST_AWARD,
    },
    "公開取得報價或企劃書": {
        "desc": "未達公告金額、逾公告金額十分之一：公開取得三家以上廠商書面報價或企劃書",
        "stages": PRE_TENDER + ["announce", "bid_opening", "evaluation", "negotiation", "award"] + POST_AWARD,
    },
    "小額採購": {
        "desc": "公告金額十分之一以下：簽准後逕洽廠商",
        "stages": ["intake", "judge", "memo", "award", "contract", "acceptance", "closeout"],
    },
    "待判斷": {
        "desc": "尚未決定招標方式：先做收件與判斷，判斷確認後以 replan 載入其餘階段",
        "stages": ["intake", "judge"],
    },
}

PROFILE_ALIASES = {
    "最有利標": "公開招標_最有利標",
    "最低標": "公開招標_最低標", "公開招標": "公開招標_最低標",
    "公開評選": "限制性招標_公開評選", "評選": "限制性招標_公開評選",
    "議價": "限制性招標_議價", "限制性招標": "限制性招標_議價",
    "公開取得": "公開取得報價或企劃書", "書面報價": "公開取得報價或企劃書", "企劃書": "公開取得報價或企劃書",
    "小額": "小額採購",
    "未定": "待判斷", "": "待判斷",
}

# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------

def now() -> str:
    return _dt.datetime.now().strftime("%Y-%m-%d %H:%M")


def roc_date(d: _dt.date | None = None) -> str:
    d = d or _dt.date.today()
    return f"{d.year - 1911}年{d.month}月{d.day}日"


def die(msg: str, code: int = 1):
    print(f"[procure] 錯誤：{msg}", file=sys.stderr)
    sys.exit(code)


def load_config() -> dict:
    if CONFIG_FILE.exists():
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    return {}


def safe_id(cid: str) -> str:
    cid = cid.strip()
    if not cid:
        die("案號不可空白")
    if re.search(r"[\\/:*?\"<>|\s]", cid):
        die("案號不可包含空白或 \\ / : * ? \" < > |")
    return cid


def case_dir(cid: str) -> Path:
    return CASES_DIR / cid


def case_file(cid: str) -> Path:
    return case_dir(cid) / "case.json"


def active_case_id() -> str | None:
    if ACTIVE_FILE.exists():
        n = ACTIVE_FILE.read_text(encoding="utf-8").strip()
        if n and case_file(n).exists():
            return n
    return None


def set_active(cid: str):
    CASES_DIR.mkdir(parents=True, exist_ok=True)
    ACTIVE_FILE.write_text(cid, encoding="utf-8")


def load_case(cid: str | None = None) -> dict:
    cid = cid or active_case_id()
    if not cid:
        die("沒有作用中的案件。請 `procure.py new <案號>` 或 `procure.py use <案號>`")
    f = case_file(cid)
    if not f.exists():
        die(f"找不到案件「{cid}」（{f}）")
    return json.loads(f.read_text(encoding="utf-8"))


def save_case(case: dict):
    f = case_file(case["id"])
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_text(json.dumps(case, ensure_ascii=False, indent=2), encoding="utf-8")
    write_state_md(case)


def make_stage(key: str, title: str | None = None, seq: int | None = None) -> dict:
    spec = STAGE_CATALOG[key]
    folder = spec["folder"] + (f"-{seq}" if seq is not None else "")
    return {
        "key": key, "name": title or spec["name"], "folder": folder,
        "step": spec["step"], "templates": list(spec["templates"]),
        "desc": spec["desc"], "confirm": spec["confirm"],
        "status": "pending", "adhoc": bool(spec.get("adhoc")),
        "started_at": None, "submitted_at": None, "confirmed_at": None,
        "revisions": [], "notes": [], "decisions": [],
    }


def log(case: dict, msg: str):
    case.setdefault("log", []).append({"time": now(), "msg": msg})


def current_index(case: dict) -> int | None:
    for i, s in enumerate(case["stages"]):
        if s["status"] in ("pending", "in_progress", "awaiting_confirm"):
            return i
    return None


def current_stage(case: dict) -> dict | None:
    i = current_index(case)
    return None if i is None else case["stages"][i]


def stage_path(case: dict, stage: dict) -> Path:
    return case_dir(case["id"]) / stage["folder"]


def fill_meta(text: str, case: dict) -> str:
    """把範本中 {{meta.xxx}} 代入案件基本資料；未填者保留佔位符供 Claude 處理"""
    meta = case.get("meta", {})
    def rep(m):
        key = m.group(1)
        if key == "案號":
            return case["id"]
        if key == "今日":
            return roc_date()
        v = meta.get(key)
        return str(v) if v not in (None, "") else m.group(0)
    return re.sub(r"\{\{meta\.([^}]+)\}\}", rep, text)


def copy_templates(case: dict, stage: dict) -> list[Path]:
    out = stage_path(case, stage)
    out.mkdir(parents=True, exist_ok=True)
    created = []
    for t in stage["templates"]:
        src = TEMPLATES_DIR / f"{t}.md"
        dst = out / f"{t}.md"
        if dst.exists():
            continue
        if src.exists():
            dst.write_text(fill_meta(src.read_text(encoding="utf-8"), case), encoding="utf-8")
        else:
            dst.write_text(f"# {t}\n\n（找不到範本 templates/{t}.md，請自行擬稿）\n", encoding="utf-8")
        created.append(dst)
    return created


def write_state_md(case: dict):
    """人類可讀的進度檔，和 case.json 同步（照原架構的 STATE.md 設計）"""
    idx = current_index(case)
    cur = None if idx is None else case["stages"][idx]
    m = case["meta"]
    lines = [f"# 案件進度｜{case['id']}　{m.get('案名', '')}", ""]
    lines.append(f"- 流程：{case['profile']}")
    lines.append(f"- 目前步驟：{'全部完成' if cur is None else f'{idx + 1}. {cur['name']}（{STATUS_LABEL[cur['status']]}）'}")
    lines.append(f"- 最後更新：{now()}")
    lines += ["", "## 基本資料", ""]
    for k, v in m.items():
        lines.append(f"- {k}：{v if v not in (None, '') else '【待補】'}")
    lines += ["", "## 階段", "", "| # | 階段 | 狀態 | 確認時間 | 修改次數 |", "|---|---|---|---|---|"]
    for i, s in enumerate(case["stages"]):
        mark = "▶ " if i == idx else ""
        lines.append(f"| {i + 1} | {mark}{s['name']}{'（臨時）' if s.get('adhoc') else ''} | {STATUS_LABEL[s['status']]} | {s['confirmed_at'] or ''} | {len(s['revisions'])} |")
    lines += ["", "## 已確認事項與使用者決定", ""]
    any_dec = False
    for s in case["stages"]:
        for d in s["decisions"]:
            lines.append(f"- [{s['name']}] {d['time']}：{d['text']}")
            any_dec = True
    if not any_dec:
        lines.append("（空）")
    lines += ["", "## 修改意見紀錄", ""]
    any_rev = False
    for s in case["stages"]:
        for r in s["revisions"]:
            lines.append(f"- [{s['name']}] {r['time']}：{r['comment']}")
            any_rev = True
    if not any_rev:
        lines.append("（空）")
    lines += ["", "## 待補檔案／待辦（承辦人自行處理）", ""]
    todos = case.get("todos", [])
    lines += [f"- {t}" for t in todos] or ["（空）"]
    lines += ["", "## 歷程", ""]
    lines += [f"- {e['time']} {e['msg']}" for e in case.get("log", [])[-30:]]
    (case_dir(case["id"]) / "STATE.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def grade_by_budget(budget, ptype) -> str:
    try:
        amt = float(str(budget).replace(",", "").replace("元", ""))
    except (ValueError, TypeError):
        return ""
    ptype = (ptype or "").strip()
    huge = {"工程": 2e8, "財物": 1e8, "勞務": 2e7}.get(ptype)
    check = {"工程": 5e7, "財物": 5e7, "勞務": 1e7}.get(ptype)
    if huge and amt >= huge:
        return "巨額採購"
    if check and amt >= check:
        return "查核金額以上未達巨額"
    if amt >= 1_500_000:
        return "公告金額以上未達查核金額"
    if amt > 150_000:
        return "未達公告金額（逾公告金額十分之一）"
    return "小額採購（公告金額十分之一以下）"


# ---------------------------------------------------------------------------
# 指令
# ---------------------------------------------------------------------------

def cmd_new(a):
    cid = safe_id(a.id)
    if case_file(cid).exists() and not a.force:
        die(f"案件「{cid}」已存在；要重建請加 --force（會保留資料夾內檔案，只重設進度）")
    cfg = load_config()
    profile = PROFILE_ALIASES.get(a.method or "", a.method or "待判斷")
    if a.stages:
        keys = [k.strip() for k in a.stages.split(",") if k.strip()]
        for k in keys:
            if k not in STAGE_CATALOG:
                die(f"未知階段 {k}；可用：{', '.join(STAGE_CATALOG)}")
        profile = "自訂"
    else:
        if profile not in PROFILES:
            die(f"未知招標方式「{a.method}」；可用：{', '.join(PROFILES)}")
        keys = PROFILES[profile]["stages"]
    meta = {
        "案名": a.title or "",
        "機關名稱": cfg.get("機關名稱", ""),
        "採購單位": cfg.get("採購單位", ""),
        "業務單位": a.unit or "",
        "採購類別": a.type or cfg.get("預設採購類別", ""),
        "預算金額": a.budget or "",
        "採購等級": grade_by_budget(a.budget, a.type or cfg.get("預設採購類別", "")) if a.budget else "",
        "招標方式": profile if profile not in ("待判斷", "自訂") else "",
        "決標方式": a.award or "",
        "承辦人": a.officer or cfg.get("承辦人", ""),
        "承辦分機": cfg.get("承辦分機", ""),
        "主持人": cfg.get("主持人", ""),
    }
    case = {"id": cid, "created_at": now(), "profile": profile, "meta": meta,
            "stages": [make_stage(k) for k in keys], "todos": [], "log": []}
    log(case, f"建立案件（流程：{profile}）")
    d = case_dir(cid)
    for sub in ("00-input",):
        (d / sub).mkdir(parents=True, exist_ok=True)
    readme = d / "00-input" / "README.md"
    if not readme.exists():
        readme.write_text(
            "把業務單位的簽、招標文件、預算書、圖說、來函等「原件複本」全部放在這個資料夾。\n"
            "流程不會改動這裡的檔案；轉出的文字會放到 01-intake/extract/。\n", encoding="utf-8")
    save_case(case)
    set_active(cid)
    print(f"已建立案件「{cid}」，流程：{profile}")
    print_status(case)


def cmd_replan(a):
    """判斷步驟確認後，依確定的招標方式重排尚未開始的階段；已確認／已跳過的保留"""
    case = load_case(a.id)
    profile = PROFILE_ALIASES.get(a.method, a.method)
    if profile not in PROFILES:
        die(f"未知招標方式「{a.method}」；可用：{', '.join(PROFILES)}")
    # 已確認、已跳過、作業中、待確認、臨時階段都保留；只重建「未開始」的階段
    kept = [s for s in case["stages"] if s["status"] != "pending" or s.get("adhoc")]
    done_keys = [s["key"] for s in kept if not s.get("adhoc")]
    kept_keys = {s["key"] for s in kept}
    new_stages = kept + [make_stage(k) for k in PROFILES[profile]["stages"] if k not in kept_keys]
    case["stages"] = new_stages
    case["profile"] = profile
    case["meta"]["招標方式"] = profile
    log(case, f"重排流程為「{profile}」（保留已確認：{', '.join(done_keys) or '無'}）")
    save_case(case)
    print(f"已依「{profile}」重排後續階段。")
    print_status(case)


def print_status(case: dict, as_json: bool = False):
    idx = current_index(case)
    if as_json:
        out = dict(case)
        out["current_index"] = idx
        out["current"] = None if idx is None else case["stages"][idx]
        out["current_dir"] = None if idx is None else str(stage_path(case, case["stages"][idx]))
        out["current_step_file"] = None if idx is None else str(STEPS_DIR / case["stages"][idx]["step"])
        out["case_dir"] = str(case_dir(case["id"]))
        out["input_dir"] = str(case_dir(case["id"]) / "00-input")
        out["templates_dir"] = str(TEMPLATES_DIR)
        out["references_dir"] = str(REFERENCE_DIR)
        out["config"] = load_config()
        print(json.dumps(out, ensure_ascii=False, indent=2))
        return
    m = case["meta"]
    print(f"案件：{case['id']}　{m.get('案名') or ''}　流程：{case['profile']}")
    print(f"類別：{m.get('採購類別') or '—'}　預算：{m.get('預算金額') or '—'}　級距：{m.get('採購等級') or '—'}")
    for i, s in enumerate(case["stages"]):
        mark = "▶" if i == idx else " "
        tag = "（臨時）" if s.get("adhoc") else ""
        rev = f"　修改 {len(s['revisions'])} 次" if s["revisions"] else ""
        print(f" {mark} {i + 1:2d}. {s['name']}{tag}　{STATUS_LABEL[s['status']]}{rev}")
    if idx is None:
        print("所有階段已完成。")
        return
    s = case["stages"][idx]
    print(f"\n目前階段：{s['name']}（{STATUS_LABEL[s['status']]}）")
    print(f"資料夾：{stage_path(case, s)}")
    print(f"作業指示：{STEPS_DIR / s['step']}")
    if s["status"] == "awaiting_confirm":
        print(f"→ 等待承辦人確認：{s['confirm']}。確認 `confirm`／修改 `revise \"意見\"`／跳過 `skip --reason`")
    elif s["status"] == "pending":
        print("→ 尚未開始。`start` 開始作業")
    else:
        print("→ 作業中。完成後 `submit`")


def cmd_status(a):
    print_status(load_case(a.id), a.json)


def cmd_brief(a):
    cid = active_case_id()
    if not cid:
        return
    case = load_case(cid)
    s = current_stage(case)
    if s is None:
        print(f"[採購流程] 案件 {cid} 所有階段已完成。")
    else:
        print(f"[採購流程] 進行中案件 {cid}（{case['meta'].get('案名', '')}）：{s['name']}（{STATUS_LABEL[s['status']]}）。輸入 /bid-run {cid} 繼續。")


def cmd_list(a):
    CASES_DIR.mkdir(parents=True, exist_ok=True)
    active = active_case_id()
    rows = []
    for d in sorted(CASES_DIR.iterdir()):
        if (d / "case.json").exists():
            c = load_case(d.name)
            s = current_stage(c)
            rows.append((d.name, c["meta"].get("案名", ""), c["profile"],
                         s["name"] if s else "已完成", STATUS_LABEL[s["status"]] if s else ""))
    if not rows:
        print("尚無案件。")
        return
    for cid, title, p, st, lb in rows:
        print(f"{'*' if cid == active else ' '} {cid}　{title}　[{p}]　{st} {lb}")
    print("（* 為作用中案件）")


def cmd_use(a):
    cid = safe_id(a.id)
    if not case_file(cid).exists():
        die(f"找不到案件「{cid}」")
    set_active(cid)
    print(f"已切換至案件「{cid}」")
    print_status(load_case(cid))


def cmd_start(a):
    case = load_case(a.id)
    s = current_stage(case)
    if s is None:
        die("所有階段已完成")
    if s["status"] == "awaiting_confirm":
        die(f"階段「{s['name']}」正等待確認，請先 confirm 或 revise")
    if s["status"] == "pending":
        s["status"] = "in_progress"
        s["started_at"] = now()
        log(case, f"開始：{s['name']}")
    files = copy_templates(case, s)
    save_case(case)
    print(f"階段「{s['name']}」作業中。資料夾：{stage_path(case, s)}")
    print(f"作業指示：{STEPS_DIR / s['step']}")
    for f in files:
        print(f"  已建立範本：{f.name}")
    if a.json:
        print_status(case, True)


def cmd_submit(a):
    case = load_case(a.id)
    s = current_stage(case)
    if s is None:
        die("所有階段已完成")
    if s["status"] == "pending":
        s["status"] = "in_progress"
        s["started_at"] = now()
    s["status"] = "awaiting_confirm"
    s["submitted_at"] = now()
    if a.note:
        s["notes"].append({"time": now(), "note": a.note})
    log(case, f"送請確認：{s['name']}")
    save_case(case)
    print(f"階段「{s['name']}」已送請承辦人確認：{s['confirm']}")
    print("Claude 必須在此結束回合，等承辦人回覆「確認」「修改…」或「跳過」。不得自動進入下一步。")


def cmd_confirm(a):
    case = load_case(a.id)
    s = current_stage(case)
    if s is None:
        die("所有階段已完成")
    if s["status"] not in ("awaiting_confirm", "in_progress") and not a.force:
        die(f"階段「{s['name']}」狀態為 {STATUS_LABEL[s['status']]}，無法確認（可加 --force）")
    s["status"] = "confirmed"
    s["confirmed_at"] = now()
    if a.note:
        s["decisions"].append({"time": now(), "text": a.note})
    log(case, f"承辦人確認：{s['name']}")
    save_case(case)
    nxt = current_stage(case)
    print(f"已確認「{s['name']}」。")
    if nxt is None:
        print("所有階段已完成。可 `export` 輸出 Word。")
    else:
        print(f"下一階段：{nxt['name']} — {nxt['desc']}")


def cmd_revise(a):
    case = load_case(a.id)
    s = current_stage(case)
    if s is None:
        die("所有階段已完成")
    s["revisions"].append({"time": now(), "comment": a.comment})
    s["status"] = "in_progress"
    log(case, f"退回修改：{s['name']}：{a.comment}")
    save_case(case)
    print(f"已記錄修改意見，「{s['name']}」回到作業中：{a.comment}")


def cmd_skip(a):
    case = load_case(a.id)
    s = current_stage(case)
    if s is None:
        die("所有階段已完成")
    s["status"] = "skipped"
    s["notes"].append({"time": now(), "note": f"跳過：{a.reason or '未註明理由'}"})
    log(case, f"跳過：{s['name']}（{a.reason or '未註明理由'}）")
    save_case(case)
    nxt = current_stage(case)
    print(f"已跳過「{s['name']}」。" + (f"下一階段：{nxt['name']}" if nxt else "所有階段已完成。"))


def cmd_add(a):
    case = load_case(a.id)
    key = ADHOC_ALIASES.get(a.kind, a.kind)
    if key not in STAGE_CATALOG:
        die(f"未知臨時階段「{a.kind}」；可用：契變／後續擴充／分期驗收")
    seq = sum(1 for s in case["stages"] if s["key"] == key) + 1
    title = a.title or f"{STAGE_CATALOG[key]['name']}（第{seq}次）"
    st = make_stage(key, title=title, seq=seq)
    idx = current_index(case)
    if idx is None:
        pos = len(case["stages"])
    elif case["stages"][idx]["status"] in ("in_progress", "awaiting_confirm") and not a.now:
        pos = idx + 1
    else:
        pos = idx
    case["stages"].insert(pos, st)
    log(case, f"插入臨時階段：{title}（第 {pos + 1} 步）")
    save_case(case)
    print(f"已插入「{title}」於第 {pos + 1} 步。")
    print_status(case)


def cmd_set(a):
    case = load_case(a.id)
    case["meta"][a.key] = a.value
    if a.key in ("預算金額", "採購類別"):
        case["meta"]["採購等級"] = grade_by_budget(case["meta"].get("預算金額"), case["meta"].get("採購類別"))
    log(case, f"基本資料：{a.key} = {a.value}")
    save_case(case)
    print(f"已更新 {a.key} = {a.value}")


def cmd_decide(a):
    """記錄使用者在目前階段做的決定（STATE.md「已確認事項與使用者決定」）"""
    case = load_case(a.id)
    s = current_stage(case)
    if s is None:
        log(case, f"決定：{a.text}")
    else:
        s["decisions"].append({"time": now(), "text": a.text})
    save_case(case)
    print("已記錄決定。")


def cmd_todo(a):
    case = load_case(a.id)
    if a.remove is not None:
        try:
            removed = case["todos"].pop(a.remove - 1)
            print(f"已移除：{removed}")
        except IndexError:
            die("沒有這個待辦編號")
    else:
        case["todos"].append(a.text)
        print("已加入待辦。")
    save_case(case)


def cmd_profiles(a):
    print("可用流程（招標方式）：")
    for k, v in PROFILES.items():
        names = " → ".join(STAGE_CATALOG[s]["name"] for s in v["stages"])
        print(f"\n■ {k}\n  {v['desc']}\n  {names}")
    print("\n階段代碼（--stages 自訂用）：")
    for k, v in STAGE_CATALOG.items():
        print(f"  {k:20s} {v['folder']:24s} {v['name']}{'（臨時）' if v.get('adhoc') else ''}")


# ---------------------------------------------------------------------------
# 匯出 Word
# ---------------------------------------------------------------------------

def md_to_docx(src: Path, dst: Path) -> str:
    try:
        import docx  # noqa: F401
        _md_to_docx_pydocx(src, dst)
        return "python-docx"
    except ImportError:
        pass
    if shutil.which("pandoc"):
        subprocess.run(["pandoc", str(src), "-o", str(dst)], check=True)
        return "pandoc"
    raise RuntimeError("沒有 python-docx 也沒有 pandoc；請 `pip install python-docx` 或安裝 pandoc")


def _md_to_docx_pydocx(src: Path, dst: Path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm

    doc = docx.Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin = sec.bottom_margin = Cm(2.5)

    def set_font(run, size=14, bold=False):
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.font.bold = bold
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), "標楷體")
        rfonts.set(qn("w:ascii"), "Times New Roman")
        rfonts.set(qn("w:hAnsi"), "Times New Roman")

    def add_par(text, size=14, bold=False, align=None, indent=0.0):
        p = doc.add_paragraph()
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        for part in re.split(r"(\*\*[^*]+\*\*)", text):
            if part.startswith("**") and part.endswith("**"):
                set_font(p.add_run(part[2:-2]), size, True)
            elif part:
                set_font(p.add_run(part), size, bold)
        p.paragraph_format.space_after = Pt(4)

    table_buf: list[str] = []

    def flush_table():
        rows = [r for r in table_buf if not re.match(r"^\s*\|?\s*:?-{2,}", r)]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        if cells:
            ncol = max(len(r) for r in cells)
            t = doc.add_table(rows=len(cells), cols=ncol)
            t.style = "Table Grid"
            for r_i, row in enumerate(cells):
                for c_i in range(ncol):
                    cell = t.cell(r_i, c_i)
                    cell.text = ""
                    set_font(cell.paragraphs[0].add_run(row[c_i] if c_i < len(row) else ""), 12, r_i == 0)
            doc.add_paragraph()
        table_buf.clear()

    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.lstrip().startswith("<!--"):
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        if table_buf:
            flush_table()
        if not line.strip():
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            add_par(m.group(2), size={1: 20, 2: 16}.get(level, 14), bold=True, align="center" if level == 1 else None)
        elif re.match(r"^\s*[-*]\s+", line):
            add_par("•" + re.sub(r"^\s*[-*]\s+", " ", line), indent=0.75)
        elif line.strip() == "---":
            doc.add_paragraph()
        else:
            add_par(line.replace("\\", ""))
        i += 1
    if table_buf:
        flush_table()
    doc.save(dst)


def cmd_export(a):
    case = load_case(a.id)
    cd = case_dir(case["id"])
    out_root = cd / "_word"
    out_root.mkdir(exist_ok=True)
    count = 0
    for s in case["stages"]:
        if a.stage and a.stage not in (s["folder"], s["name"], s["key"]):
            continue
        sp = stage_path(case, s)
        if not sp.exists():
            continue
        for md in sorted(sp.glob("*.md")):
            dst_dir = out_root / s["folder"]
            dst_dir.mkdir(parents=True, exist_ok=True)
            dst = dst_dir / (md.stem + ".docx")
            how = md_to_docx(md, dst)
            count += 1
            print(f"  {md.relative_to(cd)} → {dst.relative_to(cd)}（{how}）")
    print(f"共輸出 {count} 個 Word 檔，位於 {out_root}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="procure.py", description="採購案件流程狀態機")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("new", help="建立新案件")
    s.add_argument("id", help="案號，例如 11506")
    s.add_argument("--title", help="案名")
    s.add_argument("--method", "-m", help="招標方式／流程範本（預設：待判斷）")
    s.add_argument("--type", "-t", help="採購類別：工程／財物／勞務")
    s.add_argument("--budget", "-b", help="預算金額（數字）")
    s.add_argument("--unit", "-u", help="業務單位")
    s.add_argument("--officer", help="承辦人")
    s.add_argument("--award", help="決標方式")
    s.add_argument("--stages", help="自訂階段序列，逗號分隔")
    s.add_argument("--force", action="store_true")
    s.set_defaults(fn=cmd_new)

    s = sub.add_parser("replan", help="依確定的招標方式重排尚未開始的階段")
    s.add_argument("method")
    s.add_argument("--id", help="案號")
    s.set_defaults(fn=cmd_replan)

    for name, fn, helptext in [("status", cmd_status, "顯示案件狀態"),
                               ("start", cmd_start, "開始目前階段（複製範本到階段資料夾）")]:
        s = sub.add_parser(name, help=helptext)
        s.add_argument("id", nargs="?", help="案號（預設作用中案件）")
        s.add_argument("--json", action="store_true")
        s.set_defaults(fn=fn)

    s = sub.add_parser("brief", help="一行摘要（給 SessionStart hook）")
    s.set_defaults(fn=cmd_brief)
    s = sub.add_parser("list", help="列出所有案件")
    s.set_defaults(fn=cmd_list)
    s = sub.add_parser("use", help="切換作用中案件")
    s.add_argument("id")
    s.set_defaults(fn=cmd_use)

    s = sub.add_parser("submit", help="目前階段完成，送請承辦人確認")
    s.add_argument("--id"); s.add_argument("--note")
    s.set_defaults(fn=cmd_submit)
    s = sub.add_parser("confirm", help="承辦人確認目前階段")
    s.add_argument("--id"); s.add_argument("--note", help="同時記錄的確認事項")
    s.add_argument("--force", action="store_true")
    s.set_defaults(fn=cmd_confirm)
    s = sub.add_parser("revise", help="承辦人退回修改")
    s.add_argument("comment"); s.add_argument("--id")
    s.set_defaults(fn=cmd_revise)
    s = sub.add_parser("skip", help="跳過目前階段")
    s.add_argument("--reason"); s.add_argument("--id")
    s.set_defaults(fn=cmd_skip)
    s = sub.add_parser("add", help="插入臨時階段：契變／後續擴充／分期驗收")
    s.add_argument("kind"); s.add_argument("--title")
    s.add_argument("--now", action="store_true", help="插在目前階段之前（立即處理）")
    s.add_argument("--id")
    s.set_defaults(fn=cmd_add)
    s = sub.add_parser("set", help="更新案件基本資料")
    s.add_argument("key"); s.add_argument("value"); s.add_argument("--id")
    s.set_defaults(fn=cmd_set)
    s = sub.add_parser("decide", help="記錄使用者決定／已確認事項")
    s.add_argument("text"); s.add_argument("--id")
    s.set_defaults(fn=cmd_decide)
    s = sub.add_parser("todo", help="加入或移除承辦人待辦")
    s.add_argument("text", nargs="?", default="")
    s.add_argument("--remove", type=int, help="移除第 N 項")
    s.add_argument("--id")
    s.set_defaults(fn=cmd_todo)
    s = sub.add_parser("profiles", help="列出流程範本與階段代碼")
    s.set_defaults(fn=cmd_profiles)
    s = sub.add_parser("export", help="各階段 Markdown 輸出為 Word")
    s.add_argument("id", nargs="?"); s.add_argument("--stage")
    s.set_defaults(fn=cmd_export)
    return p


def main(argv=None):
    args = build_parser().parse_args(argv)
    args.fn(args)


if __name__ == "__main__":
    main()
