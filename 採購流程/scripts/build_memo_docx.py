#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把簽呈純文字轉成標楷體 A4 docx。
用法：python3 scripts/build_memo_docx.py <輸入.txt> <輸出.docx>

縮排規則（每行一段）：
  主旨／說明／擬辦 起頭        → 不縮排、粗體
  一、二、…                   → 第 1 層
  (一)(二)… 或 （一）         → 第 2 層
  １、２、… 或 1.             → 第 3 層
  (1)(2)…                     → 第 4 層
  其他                        → 依前一段層級
需要 python-docx：pip install python-docx
"""
import re
import sys
from pathlib import Path

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    sys.exit("需要 python-docx：pip install python-docx")

LEVELS = [
    (re.compile(r"^(主旨|說明|擬辦)[：:]"), 0, True),
    (re.compile(r"^[一二三四五六七八九十]+、"), 1, False),
    (re.compile(r"^[\(（][一二三四五六七八九十]+[\)）]"), 2, False),
    (re.compile(r"^([１２３４５６７８９０]+、|\d+\.)"), 3, False),
    (re.compile(r"^[\(（]\d+[\)）]"), 4, False),
]
HANGING = {0: 0.0, 1: 1.0, 2: 1.8, 3: 2.6, 4: 3.4}


def set_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "標楷體")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def main(src: str, dst: str):
    lines = Path(src).read_text(encoding="utf-8").splitlines()
    doc = docx.Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0)
    level = 0
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            doc.add_paragraph()
            continue
        bold = False
        matched = False
        for pat, lv, b in LEVELS:
            if pat.match(line.strip()):
                level, bold, matched = lv, b, True
                break
        p = doc.add_paragraph()
        indent = HANGING[level]
        pf = p.paragraph_format
        if matched and level > 0:
            pf.left_indent = Cm(indent)
            pf.first_line_indent = Cm(-0.8)
        elif not matched and level > 0:
            pf.left_indent = Cm(indent)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.3
        if line.startswith("簽　於") or line.startswith("簽 於"):
            set_font(p.add_run(line), 16, True)
        elif line.startswith("敬陳"):
            set_font(p.add_run(line), 14, True)
        else:
            set_font(p.add_run(line), 14, bold)
    doc.save(dst)
    print(f"已輸出 {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
